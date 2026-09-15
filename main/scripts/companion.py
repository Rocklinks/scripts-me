import os
import sys
import re
import csv
import zipfile
import shutil
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def set_col_widths(ws):
    ws.column_dimensions["A"].width = 5
    for col in ws.columns:
        if col[0].column == 1:
            continue
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value) if cell.value else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max_len + 2


def style_cells(ws, week_cells):
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align = Alignment(horizontal="center", vertical="center")
    green_font = Font(color="008000")
    red_font = Font(color="FF0000")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align
            cell.border = border
            if cell.row == 2:
                continue
            if cell.column > 1:
                headers = [c.value for c in ws[2]]
                if cell.column <= len(headers):
                    col_name = headers[cell.column - 1]
                    if col_name == "Login %" and isinstance(cell.value, float):
                        cell.number_format = "0%"
                    elif col_name == "Logged Days":
                        cell.font = green_font
                    elif col_name == "Not Logged Days":
                        cell.font = red_font
                    elif col_name and "Week" in str(col_name):
                        val = str(cell.value) if cell.value else ""
                        if val:
                            week_cells.append((cell.column_letter + str(cell.row), val))


def apply_login_bar(ws):
    headers = [c.value for c in ws[2]]
    if "Login %" not in headers:
        return
    col_idx = headers.index("Login %") + 1
    col_letter = get_column_letter(col_idx)
    last_row = ws.max_row
    rule = DataBarRule(
        start_type="num", start_value=0,
        end_type="num", end_value=1,
        color="00B050",
        showValue=True
    )
    ws.conditional_formatting.add(f"{col_letter}3:{col_letter}{last_row}", rule)


def inject_rich_text(xlsx_path, all_week_cells):
    tmp_path = xlsx_path + ".tmp"
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml") and "sheet" in item.filename.lower():
                    sheet_name = item.filename
                    week_cells_for_sheet = all_week_cells.get(sheet_name, [])
                    if week_cells_for_sheet:
                        tree = etree.fromstring(data)
                        for c in tree.iter("{%s}c" % NS):
                            ref = c.get("r")
                            if ref in dict(week_cells_for_sheet):
                                val = dict(week_cells_for_sheet)[ref]
                                for v in c.findall("{%s}v" % NS):
                                    c.remove(v)
                                for is_e in c.findall("{%s}is" % NS):
                                    c.remove(is_e)
                                c.set("t", "inlineStr")
                                is_elem = etree.SubElement(c, "{%s}is" % NS)
                                for ch in val:
                                    r_elem = etree.SubElement(is_elem, "{%s}r" % NS)
                                    rpr = etree.SubElement(r_elem, "{%s}rPr" % NS)
                                    color_elem = etree.SubElement(rpr, "{%s}color" % NS)
                                    color_elem.set("rgb", "FF0000" if ch == "O" else "008000")
                                    t_elem = etree.SubElement(r_elem, "{%s}t" % NS)
                                    t_elem.text = ch
                        data = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
                zout.writestr(item, data)
    shutil.move(tmp_path, xlsx_path)


def fix_headers(row):
    fixed = []
    for cell in row:
        if cell == "#":
            fixed.append("S.No")
            continue
        cell = re.sub(r"(Week\s+\d)(\d)", r"\1,\2", cell)
        cell = re.sub(r"(Current Week)(Week)", r"\1,\2", cell)
        cell = cell.replace("Login%", "Login %")
        cell = cell.replace("NotLoggedDays", "Not Logged Days")
        cell = cell.replace("LoggedDays", "Logged Days")
        fixed.append(cell)
    return fixed


def to_number(val, col_name):
    if col_name == "Login %":
        val = val.replace("%", "").strip()
        try:
            return float(val) / 100
        except ValueError:
            return val
    elif col_name in ("Logged Days", "Not Logged Days"):
        try:
            return int(val)
        except ValueError:
            return val
    return val


def add_csv(ws, fname, all_week_cells, sheet_idx):
    branch_name = os.path.splitext(fname)[0]
    heading = f"Employee Login- {branch_name.title()}"
    ws.append([heading])
    week_cells = []
    emp_count = 0
    with open(fname, "r", encoding="utf-8", errors="replace") as f:
        sno = 0
        for i, row in enumerate(csv.reader(f)):
            if i == 0:
                ws.append(fix_headers(row))
            else:
                sno += 1
                emp_count = sno
                row[0] = sno
                headers = [c.value for c in ws[2]]
                for j, cell in enumerate(row):
                    if j < len(headers) and headers[j] in ("Login %", "Logged Days", "Not Logged Days"):
                        row[j] = to_number(cell, headers[j])
                ws.append(row)
    max_col = ws.max_column
    last_col_letter = get_column_letter(max_col)
    ws.merge_cells(f"A1:{last_col_letter}1")
    header_cell = ws["A1"]
    header_cell.value = f"{heading}                                                                    {emp_count} Employees"
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    header_cell.font = Font(bold=True, color="FFFFFF", size=16)
    header_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    style_cells(ws, week_cells)
    apply_login_bar(ws)
    set_col_widths(ws)
    all_week_cells[f"xl/worksheets/sheet{sheet_idx}.xml"] = week_cells


def find_absent_managers(csv_files):
    absent_branches = []
    for fname in csv_files:
        branch_name = os.path.splitext(os.path.basename(fname))[0]
        with open(fname, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            headers = next(reader)
            week_indices = [i for i, h in enumerate(headers) if "week" in h.lower()]
            for row in reader:
                if len(row) > 1 and "branch manager" in row[1].lower():
                    week_values = [row[i] for i in week_indices if i < len(row)]
                    last_ox = None
                    for val in week_values:
                        if val and not all(c == "-" for c in val):
                            for ch in reversed(val):
                                if ch in ("O", "X"):
                                    last_ox = ch
                                    break
                    if last_ox == "O":
                        absent_branches.append(branch_name)
                    break
    return absent_branches


wb = Workbook()
wb.remove(wb.active)

all_week_cells = {}
sheet_idx = 0

csv_files = []
if len(sys.argv) > 1:
    for arg in sys.argv[1:]:
        if not os.path.isfile(arg):
            print(f"Skipped: {arg} (not found)")
            continue
        csv_files.append(arg)
else:
    for fname in sorted(os.listdir(".")):
        if fname.endswith(".csv"):
            csv_files.append(fname)

csv_files.sort(key=lambda x: os.path.splitext(os.path.basename(x))[0].lower())

for fname in csv_files:
    sheet_idx += 1
    heading = os.path.splitext(os.path.basename(fname))[0]
    ws = wb.create_sheet(title=heading)
    add_csv(ws, fname, all_week_cells, sheet_idx)
    print(f"Added: {fname} -> '{heading}'")

if len(csv_files) == 1:
    out = os.path.splitext(os.path.basename(csv_files[0]))[0].title() + ".xlsx"
else:
    out = "All_Branch_Companion.xlsx"
wb.save(out)
inject_rich_text(out, all_week_cells)
print(f"\nSaved to {out}")

absent = find_absent_managers(csv_files)
with open("1.txt", "w") as f:
    f.write("Today Companion App not loggedin managers\n")
    for i, branch in enumerate(absent, 1):
        f.write(f"{i}. {branch}\n")
print(f"Absent managers written to 1.txt ({len(absent)} found)")
